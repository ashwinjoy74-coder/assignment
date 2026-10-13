# School Timetable Generator
# This program makes timetables for students and teachers.
# It stops a class or teacher from having two lessons at the same time.

import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from ortools.sat.python import cp_model
except ImportError:
    cp_model = None


# ---------------- SCHOOL SETTINGS ----------------

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
PERIODS_PER_DAY = 7
CLASSES = range(1, 13)
DIVISIONS = ["A", "B", "C", "D"]

# Keep the CSV and this Python file in the same folder.
FOLDER = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(FOLDER, "requirements.csv")
OUTPUT_FILE = os.path.join(FOLDER, "school_timetable.xlsx")


# Read the subjects, teachers and weekly period counts from the CSV file.
def read_data():
    requirements = []
    lessons = []

    if not os.path.exists(INPUT_FILE):
        raise ValueError("requirements.csv was not found. Put it in the same folder as this Python file.")

    with open(INPUT_FILE, "r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)

        needed_columns = ["class", "division", "subject", "teacher", "periods_per_week"]
        if not reader.fieldnames or any(column not in reader.fieldnames for column in needed_columns):
            raise ValueError("The CSV must have these columns: class, division, subject, teacher, periods_per_week")

        for row_number, row in enumerate(reader, start=2):
            grade_text = (row.get("class") or "").strip()
            division = (row.get("division") or "").strip().upper()
            subject = (row.get("subject") or "").strip()
            teacher = (row.get("teacher") or "").strip()
            periods_text = (row.get("periods_per_week") or "").strip()

            # Ignore empty starter rows in the CSV template.
            if not subject and not teacher and not periods_text:
                continue

            if not grade_text or not division or not subject or not teacher or not periods_text:
                raise ValueError("Row " + str(row_number) + " is missing some information.")

            try:
                grade = int(grade_text)
                periods = int(periods_text)
            except ValueError:
                raise ValueError("Row " + str(row_number) + ": class and periods_per_week must be numbers.")

            if grade not in CLASSES:
                raise ValueError("Row " + str(row_number) + ": class must be between 1 and 12.")
            if division not in DIVISIONS:
                raise ValueError("Row " + str(row_number) + ": division must be A, B, C or D.")
            if periods < 1:
                raise ValueError("Row " + str(row_number) + ": periods_per_week must be at least 1.")

            class_name = str(grade) + "-" + division
            requirements.append([grade, division, subject, teacher, periods])

            # Make one lesson for every period requested in the CSV.
            for number in range(periods):
                lessons.append({
                    "class_name": class_name,
                    "subject": subject,
                    "teacher": teacher
                })

    if len(lessons) == 0:
        raise ValueError("The CSV has no subjects yet. Fill in the subject, teacher and weekly periods.")

    # Check that every division has at least one subject.
    classes_with_subjects = []
    for lesson in lessons:
        if lesson["class_name"] not in classes_with_subjects:
            classes_with_subjects.append(lesson["class_name"])

    missing_classes = []
    for grade in CLASSES:
        for division in DIVISIONS:
            name = str(grade) + "-" + division
            if name not in classes_with_subjects:
                missing_classes.append(name)

    if missing_classes:
        raise ValueError("Add subject information for these divisions: " + ", ".join(missing_classes))

    # Make sure that a class or teacher is not given too many weekly periods.
    class_load = {}
    teacher_load = {}
    for lesson in lessons:
        class_name = lesson["class_name"]
        teacher = lesson["teacher"]
        class_load[class_name] = class_load.get(class_name, 0) + 1
        teacher_load[teacher] = teacher_load.get(teacher, 0) + 1

    slots_per_week = len(DAYS) * PERIODS_PER_DAY
    for class_name in class_load:
        if class_load[class_name] > slots_per_week:
            raise ValueError(class_name + " has more lessons than the available weekly periods.")

    for teacher in teacher_load:
        if teacher_load[teacher] > slots_per_week:
            raise ValueError(teacher + " has more lessons than the available weekly periods.")

    return requirements, lessons


# Find a period for every lesson without creating class or teacher clashes.
def make_timetable(lessons):
    if cp_model is None:
        raise RuntimeError("OR-Tools is not installed. Install it with: python -m pip install -r requirements.txt")

    total_slots = len(DAYS) * PERIODS_PER_DAY
    model = cp_model.CpModel()
    choices = []

    # A lesson gets one choice for the day and period.
    for lesson_number in range(len(lessons)):
        lesson_choices = []
        for slot in range(total_slots):
            lesson_choices.append(model.NewBoolVar("lesson_" + str(lesson_number) + "_slot_" + str(slot)))
        model.AddExactlyOne(lesson_choices)
        choices.append(lesson_choices)

    # Collect lesson numbers for each class and each teacher.
    class_lessons = {}
    teacher_lessons = {}

    for lesson_number, lesson in enumerate(lessons):
        class_name = lesson["class_name"]
        teacher = lesson["teacher"]

        if class_name not in class_lessons:
            class_lessons[class_name] = []
        class_lessons[class_name].append(lesson_number)

        if teacher not in teacher_lessons:
            teacher_lessons[teacher] = []
        teacher_lessons[teacher].append(lesson_number)

    # A class can only have one lesson in a period.
    for class_name in class_lessons:
        for slot in range(total_slots):
            model.AddAtMostOne(choices[number][slot] for number in class_lessons[class_name])

    # A teacher can only teach one class in a period.
    for teacher in teacher_lessons:
        for slot in range(total_slots):
            model.AddAtMostOne(choices[number][slot] for number in teacher_lessons[teacher])

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 180
    solver.parameters.num_search_workers = 8
    result = solver.Solve(model)

    if result != cp_model.OPTIMAL and result != cp_model.FEASIBLE:
        raise RuntimeError("A timetable could not be made with these requirements. Check the teacher assignments and weekly periods.")

    # Save the chosen period number for every lesson.
    chosen_slots = []
    for lesson_number in range(len(lessons)):
        for slot in range(total_slots):
            if solver.Value(choices[lesson_number][slot]) == 1:
                chosen_slots.append(slot)
                break

    return chosen_slots


# Make the Excel workbook with one sheet for every class division and teacher.
def save_excel(requirements, lessons, chosen_slots):
    class_names = []
    teacher_names = []

    for grade in CLASSES:
        for division in DIVISIONS:
            class_names.append(str(grade) + "-" + division)

    for lesson in lessons:
        if lesson["teacher"] not in teacher_names:
            teacher_names.append(lesson["teacher"])
    teacher_names.sort()

    # Each timetable is a table with five rows (days) and seven columns (periods).
    class_tables = {}
    teacher_tables = {}

    for class_name in class_names:
        class_tables[class_name] = [["" for period in range(PERIODS_PER_DAY)] for day in DAYS]

    for teacher in teacher_names:
        teacher_tables[teacher] = [["" for period in range(PERIODS_PER_DAY)] for day in DAYS]

    # Fill the timetable tables using the periods found by the solver.
    for lesson_number, lesson in enumerate(lessons):
        slot = chosen_slots[lesson_number]
        day_number = slot // PERIODS_PER_DAY
        period_number = slot % PERIODS_PER_DAY
        class_name = lesson["class_name"]
        teacher = lesson["teacher"]

        # Check again so clashes are not accidentally written to Excel.
        if class_tables[class_name][day_number][period_number] != "":
            raise RuntimeError("A class clash was found while creating the Excel file.")
        if teacher_tables[teacher][day_number][period_number] != "":
            raise RuntimeError("A teacher clash was found while creating the Excel file.")

        class_tables[class_name][day_number][period_number] = lesson["subject"] + "\n" + teacher
        teacher_tables[teacher][day_number][period_number] = class_name + "\n" + lesson["subject"]

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "School Timetable"
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="244062")
    summary.append(["School days", len(DAYS)])
    summary.append(["Periods per day", PERIODS_PER_DAY])
    summary.append(["Number of class divisions", len(class_names)])
    summary.append(["Number of teachers", len(teacher_names)])
    summary.append(["Lessons scheduled", len(lessons)])
    summary.append(["Clash check", "No class or teacher is double-booked."])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 65

    # Include the input data in the workbook too.
    data_sheet = workbook.create_sheet("Requirements")
    data_sheet.append(["class", "division", "subject", "teacher", "periods_per_week"])
    for row in requirements:
        data_sheet.append(row)

    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")

    for column, width in zip(["A", "B", "C", "D", "E"], [12, 12, 24, 28, 20]):
        data_sheet.column_dimensions[column].width = width

    # Add one sheet per class division.
    for class_name in class_names:
        sheet = workbook.create_sheet(class_name)
        add_timetable_to_sheet(sheet, "Class " + class_name, class_tables[class_name])

    # Excel sheet names have a 31-character limit, so use T001, T002, etc.
    # The teacher's full name is shown as the heading inside the sheet.
    teacher_sheet_names = {}
    for number, teacher in enumerate(teacher_names, start=1):
        sheet_name = "T" + str(number).zfill(3)
        teacher_sheet_names[teacher] = sheet_name
        sheet = workbook.create_sheet(sheet_name)
        add_timetable_to_sheet(sheet, "Teacher: " + teacher, teacher_tables[teacher])

    # Add a list of teacher names and their worksheet names to the summary.
    summary["D1"] = "Teacher"
    summary["E1"] = "Sheet"
    summary["D1"].font = Font(bold=True, color="FFFFFF")
    summary["E1"].font = Font(bold=True, color="FFFFFF")
    summary["D1"].fill = PatternFill("solid", fgColor="244062")
    summary["E1"].fill = PatternFill("solid", fgColor="244062")
    for row_number, teacher in enumerate(teacher_names, start=2):
        summary.cell(row=row_number, column=4, value=teacher)
        summary.cell(row=row_number, column=5, value=teacher_sheet_names[teacher])
    summary.column_dimensions["D"].width = 32
    summary.column_dimensions["E"].width = 12

    workbook.save(OUTPUT_FILE)


# Format the timetable so it is easy to read in Excel.
def add_timetable_to_sheet(sheet, title, timetable):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=PERIODS_PER_DAY + 1)
    sheet.cell(row=1, column=1, value=title)
    sheet.cell(row=1, column=1).font = Font(size=15, bold=True, color="FFFFFF")
    sheet.cell(row=1, column=1).fill = PatternFill("solid", fgColor="244062")
    sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    sheet.cell(row=2, column=1, value="Day")
    for period in range(PERIODS_PER_DAY):
        sheet.cell(row=2, column=period + 2, value="Period " + str(period + 1))

    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for day_number, day in enumerate(DAYS):
        sheet.cell(row=day_number + 3, column=1, value=day)
        sheet.cell(row=day_number + 3, column=1).font = Font(bold=True)
        for period in range(PERIODS_PER_DAY):
            cell = sheet.cell(row=day_number + 3, column=period + 2, value=timetable[day_number][period])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.column_dimensions["A"].width = 15
    for column in range(2, PERIODS_PER_DAY + 2):
        sheet.column_dimensions[chr(64 + column)].width = 18
    for row in range(3, len(DAYS) + 3):
        sheet.row_dimensions[row].height = 48

    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False


def main():
    try:
        requirements, lessons = read_data()
        chosen_slots = make_timetable(lessons)
        save_excel(requirements, lessons, chosen_slots)
        print("Timetable created successfully!")
        print("Saved to:", OUTPUT_FILE)
        print("Class divisions:", len(CLASSES) * len(DIVISIONS))
        print("Teachers:", len(set(lesson["teacher"] for lesson in lessons)))
        print("Lessons scheduled:", len(lessons))
    except (ValueError, RuntimeError, OSError) as error:
        print("ERROR:", error)
        sys.exit(1)


if __name__ == "__main__":
    main()
